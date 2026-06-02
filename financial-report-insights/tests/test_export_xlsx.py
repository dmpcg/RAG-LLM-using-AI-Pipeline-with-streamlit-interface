"""Tests for Excel export module."""

import pytest

from financial_analyzer import (
    CompositeHealthScore,
    FinancialData,
    FinancialReport,
    ScenarioResult,
)


class TestFinancialExcelExporter:
    """Verify FinancialExcelExporter produces valid XLSX bytes."""

    @pytest.fixture
    def exporter(self):
        from export_xlsx import FinancialExcelExporter

        return FinancialExcelExporter()

    @pytest.fixture
    def sample_data(self):
        return FinancialData(
            total_assets=1_000_000,
            current_assets=400_000,
            cash=100_000,
            inventory=50_000,
            accounts_receivable=80_000,
            total_liabilities=500_000,
            current_liabilities=200_000,
            total_debt=300_000,
            total_equity=500_000,
            revenue=2_000_000,
            cogs=1_200_000,
            gross_profit=800_000,
            operating_income=400_000,
            net_income=300_000,
            ebit=420_000,
            ebitda=500_000,
            interest_expense=20_000,
        )

    @pytest.fixture
    def sample_results(self):
        return {
            "current_ratio": 2.0,
            "quick_ratio": 1.75,
            "gross_margin": 0.40,
            "net_margin": 0.15,
            "debt_to_equity": 0.60,
            "roe": 0.60,
        }

    # ------------------------------------------------------------------
    # export_full_report
    # ------------------------------------------------------------------

    def test_export_full_report_returns_bytes(self, exporter, sample_data, sample_results):
        result = exporter.export_full_report(sample_data, sample_results)
        assert isinstance(result, bytes)
        assert len(result) > 1000  # A real XLSX with content should be at least 1KB
        # XLSX magic bytes (PK zip archive)
        assert result[:2] == b"PK"

    def test_export_full_report_with_report(self, exporter, sample_data, sample_results):
        report = FinancialReport(
            executive_summary="Test summary",
            sections={"Overview": "Test"},
        )
        result = exporter.export_full_report(sample_data, sample_results, report=report)
        assert isinstance(result, bytes)
        assert result[:2] == b"PK"

    def test_export_full_report_with_health_score(self, exporter, sample_data):
        results = {
            "current_ratio": 2.0,
            "composite_health": CompositeHealthScore(
                score=78,
                grade="B",
                component_scores={"liquidity": 80, "profitability": 75},
                interpretation="Solid financial health.",
            ),
        }
        result = exporter.export_full_report(sample_data, results)
        assert isinstance(result, bytes)
        assert result[:2] == b"PK"

    def test_export_full_report_with_empty_results(self, exporter, sample_data):
        result = exporter.export_full_report(sample_data, {})
        assert isinstance(result, bytes)
        assert result[:2] == b"PK"

    # ------------------------------------------------------------------
    # export_ratios
    # ------------------------------------------------------------------

    def test_export_ratios_returns_bytes(self, exporter):
        ratios = {"current_ratio": 2.0, "net_margin": 0.15}
        result = exporter.export_ratios(ratios)
        assert isinstance(result, bytes)
        assert result[:2] == b"PK"

    def test_export_ratios_with_company_name(self, exporter):
        ratios = {"current_ratio": 2.0}
        result = exporter.export_ratios(ratios, company_name="Acme Corp")
        assert isinstance(result, bytes)
        assert result[:2] == b"PK"

    def test_export_empty_ratios(self, exporter):
        result = exporter.export_ratios({})
        assert isinstance(result, bytes)
        assert result[:2] == b"PK"

    def test_export_ratios_with_none_values(self, exporter):
        ratios = {"current_ratio": 2.0, "bad_metric": None}
        result = exporter.export_ratios(ratios)
        assert isinstance(result, bytes)
        assert result[:2] == b"PK"

    # ------------------------------------------------------------------
    # export_scenario_comparison
    # ------------------------------------------------------------------

    def test_export_scenario_comparison(self, exporter):
        scenario = ScenarioResult(
            scenario_name="Bull Case",
            adjustments={"revenue": 1.10},
            base_ratios={"net_margin": 0.15},
            scenario_ratios={"net_margin": 0.18},
            impact_summary="Revenue +10%",
        )
        result = exporter.export_scenario_comparison([scenario])
        assert isinstance(result, bytes)
        assert result[:2] == b"PK"

    def test_export_scenario_comparison_multiple(self, exporter):
        scenarios = [
            ScenarioResult(
                scenario_name="Bull",
                adjustments={"revenue": 1.10},
                base_ratios={"net_margin": 0.15},
                scenario_ratios={"net_margin": 0.18},
                impact_summary="Up",
            ),
            ScenarioResult(
                scenario_name="Bear",
                adjustments={"revenue": 0.90},
                base_ratios={"net_margin": 0.15},
                scenario_ratios={"net_margin": 0.12},
                impact_summary="Down",
            ),
        ]
        result = exporter.export_scenario_comparison(scenarios)
        assert isinstance(result, bytes)
        assert result[:2] == b"PK"

    def test_export_scenario_comparison_empty_list(self, exporter):
        result = exporter.export_scenario_comparison([])
        assert isinstance(result, bytes)
        assert result[:2] == b"PK"

    def test_export_scenario_with_none_ratio_values(self, exporter):
        scenario = ScenarioResult(
            scenario_name="Mixed",
            adjustments={},
            base_ratios={"net_margin": 0.15, "roe": None},
            scenario_ratios={"net_margin": 0.18},
            impact_summary="Partial data",
        )
        result = exporter.export_scenario_comparison([scenario])
        assert isinstance(result, bytes)
        assert result[:2] == b"PK"

    # ------------------------------------------------------------------
    # export_scenario_comparison — WP-2 (P1-E5) row/memory cap
    # ------------------------------------------------------------------

    def _make_big_scenario(self, name, n_keys):
        """Build a scenario with n_keys ratio rows (deterministic keys)."""
        base = {f"{name}_ratio_{i:05d}": float(i) for i in range(n_keys)}
        scen = {f"{name}_ratio_{i:05d}": float(i) + 1.0 for i in range(n_keys)}
        return ScenarioResult(
            scenario_name=name,
            adjustments={},
            base_ratios=base,
            scenario_ratios=scen,
            impact_summary=f"{name} impact",
        )

    def test_scenario_export_caps_at_whole_scenario_boundary(self, exporter):
        import io

        import openpyxl

        from export_xlsx import _MAX_EXPORT_ROWS

        assert _MAX_EXPORT_ROWS == 10_000

        # Each scenario carries ~4000 ratio rows; 4 scenarios => >10k rows total.
        # Only the scenarios that fully fit under the cap should be written.
        scenarios = [self._make_big_scenario(f"S{n}", 4000) for n in range(4)]
        result = exporter.export_scenario_comparison(scenarios)
        assert isinstance(result, bytes)
        assert result[:2] == b"PK"

        # Workbook must reopen cleanly (no out-of-order / corruption).
        wb = openpyxl.load_workbook(io.BytesIO(result))
        ws = wb["Scenario Comparison"]
        col0 = [ws.cell(row=r, column=1).value for r in range(1, ws.max_row + 1)]

        # Exactly one truncation note row.
        notes = [v for v in col0 if isinstance(v, str) and "row cap" in v]
        assert len(notes) == 1, f"expected exactly one truncation note, got {notes}"
        assert "scenarios omitted" in notes[0]

        # Not all scenarios fit; some must have been omitted.
        written_names = [v for v in col0 if v in {"S0", "S1", "S2", "S3"}]
        assert len(written_names) < 4
        assert len(written_names) >= 1

        # The LAST written scenario must have its FULL key set (no partial block).
        last_name = written_names[-1]
        last_scenario = next(s for s in scenarios if s.scenario_name == last_name)
        expected_labels = {k.replace("_", " ").title() for k in last_scenario.base_ratios}
        present_labels = set(v for v in col0 if isinstance(v, str))
        missing = expected_labels - present_labels
        assert not missing, f"last scenario block is partial; missing {len(missing)} keys"

    def test_scenario_export_small_input_uncapped(self, exporter):
        import io

        import openpyxl

        scenarios = [
            self._make_big_scenario("Bull", 3),
            self._make_big_scenario("Bear", 3),
        ]
        result = exporter.export_scenario_comparison(scenarios)
        assert isinstance(result, bytes)
        assert result[:2] == b"PK"

        wb = openpyxl.load_workbook(io.BytesIO(result))
        ws = wb["Scenario Comparison"]
        col0 = [ws.cell(row=r, column=1).value for r in range(1, ws.max_row + 1)]

        # No truncation note for small input.
        notes = [v for v in col0 if isinstance(v, str) and "row cap" in v]
        assert notes == []

        # Both scenarios present with full key sets.
        assert "Bull" in col0
        assert "Bear" in col0
        for s in scenarios:
            for k in s.base_ratios:
                assert k.replace("_", " ").title() in col0

    # ------------------------------------------------------------------
    # Scoring models
    # ------------------------------------------------------------------

    def test_export_full_report_with_z_score(self, exporter, sample_data):
        from financial_analyzer import AltmanZScore

        results = {
            "altman_z_score": AltmanZScore(
                z_score=3.1,
                zone="safe",
                components={"x1": 0.1, "x2": 0.2},
                interpretation="Safe zone.",
            ),
        }
        result = exporter.export_full_report(sample_data, results)
        assert isinstance(result, bytes)
        assert result[:2] == b"PK"

    def test_export_full_report_with_f_score(self, exporter, sample_data):
        from financial_analyzer import PiotroskiFScore

        results = {
            "piotroski_f_score": PiotroskiFScore(
                score=7,
                criteria={"roa_positive": True, "ocf_positive": True},
                interpretation="Strong value.",
            ),
        }
        result = exporter.export_full_report(sample_data, results)
        assert isinstance(result, bytes)
        assert result[:2] == b"PK"


# ---------------------------------------------------------------------------
# WS-3 P0-8: current_ratio renders as multiplier (1.50x), not percent (150.00%)
# ---------------------------------------------------------------------------


class TestRatioFormatting_P0_8:
    """Regression: ratios must NOT be formatted as percentages in XLSX output."""

    def test_current_ratio_uses_ratio_format(self):
        import io

        import xlsxwriter

        from export_xlsx import _Formats

        wb = xlsxwriter.Workbook(io.BytesIO(), {"in_memory": True})
        fmt = _Formats(wb)
        chosen = fmt.value_fmt("current_ratio")
        # ratio fmt is registered separately from pct fmt
        assert chosen is fmt.ratio
        assert chosen is not fmt.pct
        wb.close()

    def test_debt_to_equity_uses_ratio_format(self):
        import io

        import xlsxwriter

        from export_xlsx import _Formats

        wb = xlsxwriter.Workbook(io.BytesIO(), {"in_memory": True})
        fmt = _Formats(wb)
        assert fmt.value_fmt("debt_to_equity") is fmt.ratio
        wb.close()

    def test_gross_margin_still_uses_percent_format(self):
        import io

        import xlsxwriter

        from export_xlsx import _Formats

        wb = xlsxwriter.Workbook(io.BytesIO(), {"in_memory": True})
        fmt = _Formats(wb)
        assert fmt.value_fmt("gross_margin") is fmt.pct
        wb.close()

    def test_revenue_still_uses_dollar_format(self):
        import io

        import xlsxwriter

        from export_xlsx import _Formats

        wb = xlsxwriter.Workbook(io.BytesIO(), {"in_memory": True})
        fmt = _Formats(wb)
        assert fmt.value_fmt("revenue") is fmt.dollar
        wb.close()

    def test_export_ratios_writes_ratio_value_unchanged(self):
        """Sanity check: export runs without exception and current_ratio=1.5 is written."""
        from export_xlsx import FinancialExcelExporter

        exporter = FinancialExcelExporter()
        out = exporter.export_ratios({"current_ratio": 1.5, "gross_margin": 0.45})
        assert isinstance(out, bytes)
        assert len(out) > 100  # non-empty xlsx


# ---------------------------------------------------------------------------
# WS-3 WP-8: surface the pre-existing silent 500-entry cap on the Ratios sheet
# (must now log a warning AND write exactly one truncation-note row).
# ---------------------------------------------------------------------------


class TestRatiosSheetCap_WP8:
    """The Ratios sheet 500-entry cap must no longer be silent."""

    @pytest.fixture
    def exporter(self):
        from export_xlsx import FinancialExcelExporter

        return FinancialExcelExporter()

    def test_ratios_sheet_over_cap_logs_and_writes_one_note(self, exporter, caplog):
        import io
        import logging

        import openpyxl

        from export_xlsx import _MAX_RATIO_ENTRIES
        from financial_analyzer import FinancialData

        assert _MAX_RATIO_ENTRIES == 500

        # >500 numeric ratio entries forces the slice to truncate.
        results = {f"metric_{i:05d}": float(i) for i in range(_MAX_RATIO_ENTRIES + 50)}
        data = FinancialData(total_assets=1000.0, total_equity=1000.0)

        with caplog.at_level(logging.WARNING, logger="export_xlsx"):
            out = exporter.export_full_report(data, results)

        assert isinstance(out, bytes)
        assert out[:2] == b"PK"

        # A warning was logged (no longer silent).
        cap_logs = [r for r in caplog.records if "cap" in r.getMessage().lower()]
        assert cap_logs, "expected a truncation warning to be logged"

        # Reopen and assert exactly one truncation note on the Ratios sheet.
        wb = openpyxl.load_workbook(io.BytesIO(out))
        ws = wb["Ratios"]
        col0 = [ws.cell(row=r, column=1).value for r in range(1, ws.max_row + 1)]
        notes = [v for v in col0 if isinstance(v, str) and "truncated" in v.lower()]
        assert len(notes) == 1, f"expected exactly one truncation note, got {notes}"

    def test_ratios_sheet_under_cap_no_note(self, exporter):
        import io

        import openpyxl

        from financial_analyzer import FinancialData

        results = {f"metric_{i:05d}": float(i) for i in range(10)}
        data = FinancialData(total_assets=1000.0, total_equity=1000.0)
        out = exporter.export_full_report(data, results)

        wb = openpyxl.load_workbook(io.BytesIO(out))
        ws = wb["Ratios"]
        col0 = [ws.cell(row=r, column=1).value for r in range(1, ws.max_row + 1)]
        notes = [v for v in col0 if isinstance(v, str) and "truncated" in v.lower()]
        assert notes == []

    def test_export_ratios_over_cap_logs_and_writes_one_note(self, exporter, caplog):
        """The single-sheet export_ratios path also caps + surfaces."""
        import io
        import logging

        import openpyxl

        from export_xlsx import _MAX_RATIO_ENTRIES

        ratios = {f"metric_{i:05d}": float(i) for i in range(_MAX_RATIO_ENTRIES + 50)}

        with caplog.at_level(logging.WARNING, logger="export_xlsx"):
            out = exporter.export_ratios(ratios)

        assert out[:2] == b"PK"
        cap_logs = [r for r in caplog.records if "cap" in r.getMessage().lower()]
        assert cap_logs, "expected a truncation warning to be logged"

        wb = openpyxl.load_workbook(io.BytesIO(out))
        ws = wb["Ratios"]
        col0 = [ws.cell(row=r, column=1).value for r in range(1, ws.max_row + 1)]
        notes = [v for v in col0 if isinstance(v, str) and "truncated" in v.lower()]
        assert len(notes) == 1, f"expected exactly one truncation note, got {notes}"


class TestFormulaInjection_AUD0522_05:
    """User-controlled strings beginning with = + - @ must be stored as literal
    strings, never interpreted as formulas (CWE-1236)."""

    def test_exporter_does_not_promote_user_string_to_formula(self):
        import io

        import openpyxl

        from export_xlsx import FinancialExcelExporter
        from financial_analyzer import ScenarioResult

        # scenario_name is attacker-controlled and written via ws.write(); a
        # leading '=' must NOT be promoted to a live formula in the workbook.
        exporter = FinancialExcelExporter()
        evil = '=HYPERLINK("http://evil","click")'
        out = exporter.export_scenario_comparison(
            [ScenarioResult(scenario_name=evil)]
        )
        assert out[:2] == b"PK"

        wb = openpyxl.load_workbook(io.BytesIO(out))
        ws = wb["Scenario Comparison"]
        cells = [ws.cell(row=r, column=1) for r in range(1, ws.max_row + 1)]
        evil_cells = [c for c in cells if c.value == evil]
        assert evil_cells, "expected the scenario name to be written verbatim"
        for c in evil_cells:
            assert c.data_type != "f", "user string must not be stored as a formula"

    def test_leading_equals_string_not_a_formula(self):
        import io

        import openpyxl
        import xlsxwriter

        buf = io.BytesIO()
        wb = xlsxwriter.Workbook(buf, {"in_memory": True, "strings_to_formulas": False})
        ws = wb.add_worksheet("S")
        ws.write(0, 0, "=cmd|'/c calc'!A1")
        wb.close()

        loaded = openpyxl.load_workbook(io.BytesIO(buf.getvalue()))
        cell = loaded["S"].cell(row=1, column=1)
        assert cell.value == "=cmd|'/c calc'!A1"
        assert cell.data_type == "s", "leading '=' string must be stored as text, not a formula"
